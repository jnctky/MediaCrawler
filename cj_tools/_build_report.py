# -*- coding: utf-8 -*-
"""
MediaCrawler 抓取结果 -> 报表（通用版）
- 自动发现 data/xhs/json 下最新的 search_contents_*.json / search_comments_*.json
- 自动从 source_keyword 推断关键词
- 兼容 "10万+" / "8.6万" / "5044" 等多种互动数格式
- 输出：Excel（多工作表）+ 结构化 JSON（供 HTML 报告使用）

用法：
  python _build_report.py                          # 自动用最新数据
  python _build_report.py --contents a.json --comments b.json --label 手账
"""
import json, collections, datetime, statistics, os, sys, io, re, glob, argparse

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

_THIS = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(_THIS)               # MediaCrawler 仓库根（cj_tools 的上一级）
OUT = os.path.join(BASE, "报告")
DATA_DIR = os.path.join(BASE, "data", "xhs", "json")


def parse_cnt(v):
    """'10万+' / '1.2万' / '5044' / 5044 -> int"""
    if v is None:
        return 0
    s = str(v).strip().replace(",", "").replace("+", "").replace(" ", "")
    if not s:
        return 0
    m = re.match(r"^([\d.]+)万$", s)
    if m:
        return int(float(m.group(1)) * 10000)
    m = re.match(r"^([\d.]+)亿$", s)
    if m:
        return int(float(m.group(1)) * 100000000)
    try:
        return int(float(s))
    except Exception:
        return 0


def load(path):
    with open(path, encoding="utf-8") as f:
        t = f.read().strip()
    if t.startswith("["):
        return json.loads(t)
    return [json.loads(ln) for ln in t.splitlines() if ln.strip()]


def latest(pattern):
    fs = [p for p in glob.glob(os.path.join(DATA_DIR, pattern))
          if not os.path.basename(p).startswith("filtered_")]
    if not fs:
        return None
    return max(fs, key=os.path.getmtime)


def ts(ms):
    if not ms:
        return ""
    return datetime.datetime.fromtimestamp(int(ms) / 1000).strftime("%Y-%m-%d %H:%M")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--contents", default=None, help="笔记 json/jsonl")
    ap.add_argument("--comments", default=None, help="评论 json/jsonl")
    ap.add_argument("--label", default=None, help="关键词标签（默认自动推断）")
    ap.add_argument("--outdir", default=OUT)
    args = ap.parse_args()

    cf = args.contents or latest("search_contents_*.json*")
    mf = args.comments or latest("search_comments_*.json*")
    if not cf:
        print("找不到笔记数据文件"); return

    notes_raw = load(cf)
    cmts_raw = load(mf) if mf and os.path.exists(mf) else []

    # 去重
    notes, seen = [], set()
    for n in notes_raw:
        if n.get("note_id") in seen:
            continue
        seen.add(n.get("note_id")); notes.append(n)
    cmts, seen2 = [], set()
    for c in cmts_raw:
        if c.get("comment_id") in seen2:
            continue
        seen2.add(c.get("comment_id")); cmts.append(c)

    kwset = []
    for n in notes:
        k = (n.get("source_keyword") or "").strip()
        if k and k not in kwset:
            kwset.append(k)
    label = args.label or ("+".join(kwset) if kwset else "未标注")
    if args.label is None and len(kwset) == 1:
        label = kwset[0]

    print(f"数据源: {os.path.basename(cf)} / {os.path.basename(mf) if mf else '(无评论)'}")
    print(f"关键词: {kwset}")
    print(f"笔记 {len(notes_raw)} -> 去重 {len(notes)} | 评论 {len(cmts_raw)} -> 去重 {len(cmts)}")

    os.makedirs(args.outdir, exist_ok=True)

    # ---------- Excel ----------
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    RED, thin = "FF2442", Side(style="thin", color="DDDDDD")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    HEAD_FILL = PatternFill("solid", fgColor=RED)
    HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    CELL_FONT = Font(name="微软雅黑", size=10)

    wb = Workbook()

    def hdr(ws):
        for c in ws[1]:
            c.fill = HEAD_FILL; c.font = HEAD_FONT
            c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORDER
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = ws.cell(row=2, column=1)

    def body(ws):
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.font = CELL_FONT; c.border = BORDER

    def fit(ws, maxw=60):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[letter].width = min(max(w * 1.2 + 3, 9), maxw)

    ws = wb.active; ws.title = "笔记(去重)"
    ws.append(["排名", "标题/正文摘要", "类型", "点赞", "收藏", "评论数", "分享",
               "总互动", "收藏/点赞", "发布时间", "来源关键词", "话题标签", "笔记链接"])
    sn = sorted(notes, key=lambda x: parse_cnt(x.get("liked_count")), reverse=True)
    for i, n in enumerate(sn, 1):
        like, col = parse_cnt(n.get("liked_count")), parse_cnt(n.get("collected_count"))
        ws.append([i, (n.get("title") or n.get("desc", "")).replace("\n", " ")[:80],
                   "视频" if n.get("type") == "video" else "图文",
                   like, col, parse_cnt(n.get("comment_count")), parse_cnt(n.get("share_count")),
                   like + col + parse_cnt(n.get("comment_count")) + parse_cnt(n.get("share_count")),
                   round(col / like, 2) if like else "",
                   ts(n.get("time")), n.get("source_keyword", ""),
                   (n.get("tag_list") or "").replace(",", " "), n.get("note_url", "")])
    hdr(ws); body(ws); fit(ws)

    ws2 = wb.create_sheet("评论(去重)")
    ws2.append(["排名", "评论内容", "点赞", "子评论数", "所属笔记ID", "昵称(已打码)", "时间"])
    for i, c in enumerate(sorted(cmts, key=lambda x: parse_cnt(x.get("like_count")), reverse=True), 1):
        ws2.append([i, (c.get("content") or "").replace("\n", " ")[:120],
                    parse_cnt(c.get("like_count")), parse_cnt(c.get("sub_comment_count")),
                    c.get("note_id", ""), c.get("nickname", ""), ts(c.get("create_time"))])
    hdr(ws2); body(ws2); fit(ws2, 70)

    ws3 = wb.create_sheet("关键词对比")
    ws3.append(["关键词", "笔记数", "最高赞", "中位赞", "均值赞", "总点赞", "破千篇数", "破万篇数", "平均收藏", "视频占比"])
    for kw in kwset:
        sub = [n for n in notes if (n.get("source_keyword") or "") == kw]
        lks = [parse_cnt(n.get("liked_count")) for n in sub]
        cls = [parse_cnt(n.get("collected_count")) for n in sub]
        vids = sum(1 for n in sub if n.get("type") == "video")
        ws3.append([kw, len(sub), max(lks) if lks else 0,
                    statistics.median(lks) if lks else 0,
                    round(statistics.mean(lks)) if lks else 0, sum(lks),
                    sum(1 for x in lks if x >= 1000), sum(1 for x in lks if x >= 10000),
                    round(statistics.mean(cls)) if cls else 0,
                    f"{vids/len(sub)*100:.0f}%" if sub else ""])
    hdr(ws3); body(ws3); fit(ws3)

    ws4 = wb.create_sheet("标签热度")
    tagc = collections.Counter()
    for n in notes:
        for t in (n.get("tag_list") or "").split(","):
            if t.strip():
                tagc[t.strip()] += 1
    ws4.append(["排名", "话题标签", "出现笔记数", "占比"])
    for i, (k, v) in enumerate(tagc.most_common(50), 1):
        ws4.append([i, "#" + k, v, f"{v/len(notes)*100:.0f}%"])
    hdr(ws4); body(ws4); fit(ws4)

    safe = re.sub(r'[\\/:*?"<>|]', "_", label)
    xlsx = os.path.join(args.outdir, f"MediaCrawler-小红书抓取报告_{safe}.xlsx")
    wb.save(xlsx)
    print("Excel:", xlsx)

    # ---------- 结构化 JSON ----------
    lks = [parse_cnt(n.get("liked_count")) for n in notes]
    tag_list = [{"name": k, "count": v} for k, v in tagc.most_common(30)]
    report = {
        "label": label, "keywords": kwset,
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source_files": [os.path.basename(cf)] + ([os.path.basename(mf)] if mf else []),
        "raw_notes": len(notes_raw), "notes": len(notes),
        "raw_cmts": len(cmts_raw), "cmts": len(cmts),
        "type_dist": dict(collections.Counter(("视频" if n.get("type") == "video" else "图文") for n in notes)),
        "like_stats": {
            "max": max(lks) if lks else 0,
            "median": statistics.median(lks) if lks else 0,
            "mean": round(statistics.mean(lks)) if lks else 0,
            "min": min(lks) if lks else 0,
            "over1000": sum(1 for x in lks if x >= 1000),
            "over5000": sum(1 for x in lks if x >= 5000),
            "over10000": sum(1 for x in lks if x >= 10000),
        },
        "top_notes": [{
            "title": (n.get("title") or n.get("desc", "")).replace("\n", " ")[:60],
            "type": "视频" if n.get("type") == "video" else "图文",
            "liked": parse_cnt(n.get("liked_count")), "collected": parse_cnt(n.get("collected_count")),
            "comment": parse_cnt(n.get("comment_count")), "share": parse_cnt(n.get("share_count")),
            "date": ts(n.get("time"))[:10], "url": n.get("note_url", ""),
            "tags": [t for t in (n.get("tag_list") or "").split(",") if t][:4],
        } for n in sn[:30]],
        "all_liked": sorted(lks, reverse=True),
        "tags": tag_list,
        "kw_stats": [], "top_cmts": [], "note_type_like": {},
        "time_span": "",
    }
    tss = sorted(int(n["time"]) for n in notes if n.get("time"))
    if tss:
        a = datetime.datetime.fromtimestamp(tss[0] / 1000); b = datetime.datetime.fromtimestamp(tss[-1] / 1000)
        report["time_span"] = f"{a:%Y-%m-%d} ~ {b:%Y-%m-%d}"
        report["time_hist"] = collections.Counter(
            f"{datetime.datetime.fromtimestamp(t/1000):%Y-%m}" for t in tss)
        report["time_hist"] = dict(sorted(report["time_hist"].items()))

    for kw in kwset:
        sub = [n for n in notes if (n.get("source_keyword") or "") == kw]
        sl = [parse_cnt(n.get("liked_count")) for n in sub]
        sc = [parse_cnt(n.get("collected_count")) for n in sub]
        report["kw_stats"].append({
            "kw": kw, "count": len(sub), "max": max(sl), "median": statistics.median(sl),
            "mean": round(statistics.mean(sl)), "total": sum(sl),
            "over1000": sum(1 for x in sl if x >= 1000),
            "over10000": sum(1 for x in sl if x >= 10000),
            "avg_col": round(statistics.mean(sc)) if sc else 0,
            "video_pct": round(sum(1 for n in sub if n.get("type") == "video") / len(sub) * 100) if sub else 0,
        })
    for t, lab in [("video", "视频"), ("normal", "图文")]:
        sub = [parse_cnt(n.get("liked_count")) for n in notes if n.get("type") == t]
        if sub:
            report["note_type_like"][lab] = {"n": len(sub), "mean": round(statistics.mean(sub)),
                                             "median": statistics.median(sub)}
    report["top_cmts"] = [{
        "content": (c.get("content") or "").replace("\n", " ")[:100],
        "like": parse_cnt(c.get("like_count")), "note": c.get("note_id", "")
    } for c in sorted(cmts, key=lambda x: parse_cnt(x.get("like_count")), reverse=True)[:10]]

    jp = os.path.join(args.outdir, f"_report_data_{safe}.json")
    with open(jp, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=1)
    print("摘要 JSON:", jp)
    print()
    print(json.dumps({k: report[k] for k in ["label", "notes", "cmts", "time_span", "like_stats", "type_dist"]},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
